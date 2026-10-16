"""Analýza útrat (jen megaAdmin AGC): ceny nových kotoučů a kolik se ušetřilo reprofilací.

Ceny nových kotoučů se zadávají podle typu kotouče (výrobce + průměr + profil) a platí pro
všechny kotouče toho typu. Úspora jedné dokončené reprofilace = cena nového kotouče − cena reprofilace.
Reprofilace bez ceny nebo kotouče bez ceny v tabulce se do úspory nepočítají — analýza je vypíše zvlášť.
"""
import io
from collections import defaultdict
from datetime import date, datetime, time, timedelta

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session, joinedload

from . import crud, models, schemas
from .database import get_db
from .deps import require_mega_admin
from .repair import RESULT, _sheet

router = APIRouter(prefix="/api/analytics", tags=["analýza útrat"])

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CZK, DT = '# ##0 "Kč"', "d.m.yyyy h:mm"


def _xlsx(wb: Workbook, name: str) -> StreamingResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="{name}"'})


# ---------- Ceny nových kotoučů ----------

def _price_out(db: Session, p: models.WheelPrice) -> schemas.WheelPriceOut:
    wheels = db.scalar(select(func.count(models.Wheel.id)).where(
        models.Wheel.manufacturer_id == p.manufacturer_id,
        models.Wheel.diameter == p.diameter,
        models.Wheel.profile_id == p.profile_id,
    )) or 0
    return schemas.WheelPriceOut(
        id=p.id, manufacturer=p.manufacturer.name, diameter=p.diameter, profile=p.profile.code,
        price=p.price, updated_at=p.updated_at, updated_by_username=p.updated_by_username, wheels=wheels,
    )


def _all_prices(db: Session) -> list[models.WheelPrice]:
    return list(db.scalars(
        select(models.WheelPrice)
        .options(joinedload(models.WheelPrice.manufacturer), joinedload(models.WheelPrice.profile))
    ).all())


def _save_price(db: Session, price: models.WheelPrice | None, data: schemas.WheelPriceIn, user: models.User) -> models.WheelPrice:
    manufacturer_name, profile_code = data.manufacturer.strip(), data.profile.strip()
    if not manufacturer_name or not profile_code:
        raise HTTPException(400, "Vyplňte výrobce i profil kotouče")
    manufacturer = crud.get_or_create_manufacturer(db, manufacturer_name)
    profile = crud.get_or_create_profile(db, profile_code)
    duplicate = db.scalar(select(models.WheelPrice.id).where(
        models.WheelPrice.manufacturer_id == manufacturer.id,
        models.WheelPrice.diameter == data.diameter,
        models.WheelPrice.profile_id == profile.id,
    ))
    if duplicate and (price is None or duplicate != price.id):
        raise HTTPException(409, f"Cena pro {manufacturer.name} · Ø {data.diameter} mm · {profile.code} už je zadaná — upravte ji v tabulce")
    if price is None:
        price = models.WheelPrice()
        db.add(price)
    price.manufacturer, price.diameter, price.profile, price.price = manufacturer, data.diameter, profile, data.price
    price.updated_at, price.updated_by_username = models.utc_now(), user.username
    db.commit()
    db.refresh(price)
    return price


@router.get("/prices", response_model=list[schemas.WheelPriceOut])
def list_prices(db: Session = Depends(get_db), _admin: models.User = Depends(require_mega_admin)):
    rows = [_price_out(db, p) for p in _all_prices(db)]
    return sorted(rows, key=lambda r: (r.manufacturer.lower(), r.diameter, r.profile.lower()))


@router.post("/prices", response_model=schemas.WheelPriceOut, status_code=201)
def create_price(data: schemas.WheelPriceIn, db: Session = Depends(get_db), admin: models.User = Depends(require_mega_admin)):
    return _price_out(db, _save_price(db, None, data, admin))


@router.put("/prices/{price_id}", response_model=schemas.WheelPriceOut)
def update_price(price_id: int, data: schemas.WheelPriceIn, db: Session = Depends(get_db),
                 admin: models.User = Depends(require_mega_admin)):
    price = db.get(models.WheelPrice, price_id)
    if not price:
        raise HTTPException(404, "Cena nenalezena")
    return _price_out(db, _save_price(db, price, data, admin))


@router.delete("/prices/{price_id}", status_code=204)
def delete_price(price_id: int, db: Session = Depends(get_db), _admin: models.User = Depends(require_mega_admin)):
    price = db.get(models.WheelPrice, price_id)
    if not price:
        raise HTTPException(404, "Cena nenalezena")
    db.delete(price)
    db.commit()
    return Response(status_code=204)


@router.get("/prices/export.xlsx")
def export_prices(db: Session = Depends(get_db), _admin: models.User = Depends(require_mega_admin)):
    rows = list_prices(db, _admin)
    wb = Workbook()
    wb.remove(wb.worksheets[0])
    _sheet(wb, "Ceny kotoučů", [
        ("Výrobce", 20, None), ("Průměr (mm)", 12, "0"), ("Profil", 10, None), ("Cena nového kotouče", 20, CZK),
        ("Kotoučů v evidenci", 18, "0"), ("Naposledy změnil", 18, None), ("Změněno", 17, DT),
    ], [[r.manufacturer, r.diameter, r.profile, r.price, r.wheels, r.updated_by_username, r.updated_at] for r in rows])
    return _xlsx(wb, f"ceny-kotoucu-{date.today():%Y-%m-%d}.xlsx")


# ---------- Analýza úspor ----------

def _reprofilings(db: Session, date_from: date | None, date_to: date | None,
                  firm_id: int | None) -> list[models.WheelReprofiling]:
    """Dokončené reprofilace (hotovo / nelze) podle data dokončení v zadaném období (volitelně jen jedna firma)."""
    q = (
        select(models.WheelReprofiling)
        .where(models.WheelReprofiling.result.is_not(None), models.WheelReprofiling.finished_at.is_not(None))
        .options(joinedload(models.WheelReprofiling.wheel).joinedload(models.Wheel.manufacturer),
                 joinedload(models.WheelReprofiling.wheel).joinedload(models.Wheel.profile))
        .order_by(models.WheelReprofiling.finished_at)
    )
    if date_from:
        q = q.where(models.WheelReprofiling.finished_at >= datetime.combine(date_from, time.min))
    if date_to:
        q = q.where(models.WheelReprofiling.finished_at < datetime.combine(date_to + timedelta(days=1), time.min))
    if firm_id is not None:
        q = q.where(models.WheelReprofiling.location_id == firm_id)
    return list(db.scalars(q).unique().all())


def _firm(db: Session, firm_id: int | None) -> models.Location | None:
    if firm_id is None:
        return None
    firm = db.get(models.Location, firm_id)
    if not firm or firm.type != models.LocationType.REPROFILACE:
        raise HTTPException(404, "Firma reprofilace nenalezena")
    return firm


def _analyse(db: Session, date_from: date | None, date_to: date | None, firm_id: int | None):
    if date_from and date_to and date_from > date_to:
        raise HTTPException(400, "Datum „od“ musí být před datem „do“")
    _firm(db, firm_id)
    prices = {(p.manufacturer_id, p.diameter, p.profile_id): p.price for p in _all_prices(db)}
    detail, missing = [], defaultdict(int)
    for r in _reprofilings(db, date_from, date_to, firm_id):
        w = r.wheel
        new_price = prices.get((w.manufacturer_id, w.diameter, w.profile_id))
        done = r.result == "hotovo"
        if done and new_price is None:
            missing[(w.manufacturer.name, w.diameter, w.profile.code)] += 1
        saving = new_price - r.cost if done and new_price is not None and r.cost is not None else None
        detail.append((r, w, new_price, saving))
    return detail, missing


@router.get("/savings", response_model=schemas.SavingsAnalysis)
def savings(date_from: date | None = Query(None), date_to: date | None = Query(None), firm_id: int | None = Query(None),
            db: Session = Depends(get_db), _admin: models.User = Depends(require_mega_admin)):
    detail, missing = _analyse(db, date_from, date_to, firm_id)
    done = [d for d in detail if d[0].result == "hotovo"]
    complete = [d for d in done if d[3] is not None]

    months: dict[str, schemas.SavingsMonth] = {}
    firm_months: dict[tuple[str, str], schemas.SavingsMonth] = {}
    for r, _w, new_price, saving in done:
        key = f"{r.finished_at:%Y-%m}"
        name = r.location_name or "Neuvedená firma"
        for m in (months.setdefault(key, schemas.SavingsMonth(month=key, saved=0, spent=0, count=0)),
                  firm_months.setdefault((key, name), schemas.SavingsMonth(month=key, firm=name, saved=0, spent=0, count=0))):
            m.count += 1
            m.spent += r.cost or 0
            m.saved += saving or 0

    firms: dict[str, dict] = {}
    for r, _w, _p, saving in detail:
        name = r.location_name or "Neuvedená firma"
        f = firms.setdefault(name, {"done": 0, "failed": 0, "spent": 0.0, "with_cost": 0, "saved": 0.0, "days": []})
        if r.result == "hotovo":
            f["done"] += 1
            if r.cost is not None:
                f["spent"] += r.cost
                f["with_cost"] += 1
            f["saved"] += saving or 0
        else:
            f["failed"] += 1
        f["days"].append((r.finished_at - r.sent_at).total_seconds() / 86400)
    firm_rows = [
        schemas.SavingsFirm(
            name=name, done=f["done"], failed=f["failed"],
            success_rate=round(100 * f["done"] / (f["done"] + f["failed"]), 1),
            spent=f["spent"], saved=f["saved"],
            avg_cost=round(f["spent"] / f["with_cost"], 2) if f["with_cost"] else None,
            avg_days=round(sum(f["days"]) / len(f["days"]), 1) if f["days"] else None,
        )
        for name, f in firms.items()
    ]
    firm_rows.sort(key=lambda x: x.saved, reverse=True)

    return schemas.SavingsAnalysis(
        saved=sum(d[3] for d in complete),
        spent=sum(d[0].cost or 0 for d in done),
        new_value=sum(d[2] for d in complete),
        done=len(done),
        failed=len(detail) - len(done),
        without_cost=sum(1 for d in done if d[0].cost is None),
        without_price=sum(missing.values()),
        months=sorted(months.values(), key=lambda m: m.month),
        firm_months=sorted(firm_months.values(), key=lambda m: (m.month, m.firm or "")),
        firms=firm_rows,
        missing_prices=[
            schemas.MissingPrice(manufacturer=k[0], diameter=k[1], profile=k[2], count=n)
            for k, n in sorted(missing.items(), key=lambda kv: -kv[1])
        ],
    )


@router.get("/savings/export.xlsx")
def export_savings(date_from: date | None = Query(None), date_to: date | None = Query(None), firm_id: int | None = Query(None),
                   db: Session = Depends(get_db), _admin: models.User = Depends(require_mega_admin)):
    data = savings(date_from, date_to, firm_id, db, _admin)
    detail, _missing = _analyse(db, date_from, date_to, firm_id)
    firm = _firm(db, firm_id)
    wb = Workbook()
    wb.remove(wb.worksheets[0])
    period = f"{date_from:%d.%m.%Y}" if date_from else "začátek evidence"
    period += f" – {date_to:%d.%m.%Y}" if date_to else " – dnes"
    _sheet(wb, "Souhrn", [("Ukazatel", 34, None), ("Hodnota", 22, None)], [
        ["Období", period],
        ["Firma reprofilace", firm.name if firm else "Všechny firmy"],
        ["Ušetřeno reprofilací", data.saved],
        ["Utraceno za reprofilace", data.spent],
        ["Nové kotouče by stály", data.new_value],
        ["Dokončených reprofilací", data.done],
        ["Nešlo reprofilovat", data.failed],
        ["Reprofilace bez ceny", data.without_cost],
        ["Reprofilace kotoučů bez ceny nového", data.without_price],
    ])
    ws = wb["Souhrn"]
    for line in (4, 5, 6):
        ws.cell(row=line, column=2).number_format = CZK
    _sheet(wb, "Firmy", [
        ("Firma", 28, None), ("Hotovo", 9, "0"), ("Nešlo", 9, "0"), ("Úspěšnost (%)", 14, "0.0"),
        ("Utraceno", 16, CZK), ("Ušetřeno", 16, CZK), ("Průměrná cena", 16, CZK), ("Průměrná doba (dny)", 19, "0.0"),
    ], [[f.name, f.done, f.failed, f.success_rate, f.spent, f.saved, f.avg_cost, f.avg_days] for f in data.firms])
    _sheet(wb, "Po měsících", [("Měsíc", 12, None), ("Reprofilací", 12, "0"), ("Utraceno", 16, CZK), ("Ušetřeno", 16, CZK)],
           [[m.month, m.count, m.spent, m.saved] for m in data.months])
    if firm is None and data.firms:
        # Porovnání firem po měsících: řádek = měsíc, pro každou firmu ušetřeno / utraceno / počet
        names = [f.name for f in data.firms]
        by_key = {(m.month, m.firm): m for m in data.firm_months}
        columns: list[tuple[str, int, str | None]] = [("Měsíc", 12, None)]
        for n in names:
            columns += [(f"{n} – ušetřeno", 22, CZK), (f"{n} – utraceno", 22, CZK), (f"{n} – počet", 14, "0")]
        rows = []
        for m in data.months:
            row: list = [m.month]
            for n in names:
                fm = by_key.get((m.month, n))
                row += [fm.saved, fm.spent, fm.count] if fm else [0, 0, 0]
            rows.append(row)
        _sheet(wb, "Po měsících podle firem", columns, rows)
    _sheet(wb, "Reprofilace", [
        ("Číslo kotouče", 16, None), ("Výrobce", 16, None), ("Průměr (mm)", 11, "0"), ("Profil", 9, None),
        ("Firma", 24, None), ("Odesláno", 17, DT), ("Dokončeno", 17, DT), ("Výsledek", 18, None),
        ("Cena reprofilace", 16, CZK), ("Cena nového", 16, CZK), ("Úspora", 16, CZK),
    ], [[w.serial_number, w.manufacturer.name, w.diameter, w.profile.code, r.location_name, r.sent_at, r.finished_at,
         RESULT.get(r.result, r.result), r.cost, new_price, saving] for r, w, new_price, saving in detail])
    if data.missing_prices:
        _sheet(wb, "Chybí cena nového", [("Výrobce", 20, None), ("Průměr (mm)", 12, "0"), ("Profil", 10, None), ("Reprofilací", 12, "0")],
               [[m.manufacturer, m.diameter, m.profile, m.count] for m in data.missing_prices])
    suffix = f"-firma-{firm.id}" if firm else ""
    return _xlsx(wb, f"analyza-utrat{suffix}-{date.today():%Y-%m-%d}.xlsx")

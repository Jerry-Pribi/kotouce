"""Část Reprofilace: API pro firmu, která pro AGC přebrušuje (reprofiluje) kotouče.

Účet reprofilace patří k jednomu místu reprofilace (Lokace typu „reprofilace“).
Vidí jen kotouče, které má právě u sebe, a historii svých reprofilací. Hlásí:
  převzetí kotouče → reprofilace hotová (poznámka, nepovinná cena) / nelze reprofilovat.
Stav kotouče v evidenci AGC se tím nemění — kotouč zpátky převezme AGC (změna stavu na kartě kotouče).
Export do Excelu (i s cenami) je jen tady, v části Reprofilace.
"""
import io
from datetime import date, datetime

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from . import crud, models, schemas
from .database import get_db
from .deps import require_repair

router = APIRouter(prefix="/api/repair", tags=["reprofilace"])
S = models.WheelStatus


def _is_at(wheel: models.Wheel, location_id: int | None) -> bool:
    """Kotouč je teď u této firmy (na reprofilaci a otevřený záznam nebo stanoviště je u ní)."""
    if wheel.status != S.NA_REPROFILACI or location_id is None:
        return False
    current = crud.open_reprofiling(wheel)
    if current and current.location_id is not None:
        return current.location_id == location_id
    return wheel.location_id == location_id


def _can_see(wheel: models.Wheel, location_id: int | None) -> bool:
    return _is_at(wheel, location_id) or any(r.location_id == location_id for r in wheel.reprofilings)


def _get_wheel(db: Session, wheel_id: int, user: models.User) -> models.Wheel:
    wheel = db.get(models.Wheel, wheel_id)
    if not wheel or not _can_see(wheel, user.location_id):
        raise HTTPException(404, "Kotouč nenalezen")
    return wheel


def _current_here(db: Session, wheel: models.Wheel, user: models.User) -> models.WheelReprofiling:
    """Rozpracovaná reprofilace u této firmy; u starších dat bez záznamu ho založí."""
    if not _is_at(wheel, user.location_id):
        raise HTTPException(400, "Kotouč teď není u vás na reprofilaci")
    current = crud.open_reprofiling(wheel)
    if current is None:
        current = models.WheelReprofiling(
            wheel_id=wheel.id, sent_at=models.utc_now(), sent_by_username="Historický záznam",
            meters_before=crud.wheel_total_meters(db, wheel.id),
        )
        db.add(current)
        wheel.reprofilings.insert(0, current)
    if current.location_id is None and user.location:
        current.location_id, current.location_name = user.location.id, user.location.name
    return current


def _to_out(db: Session, wheel: models.Wheel, user: models.User) -> schemas.RepairWheelOut:
    here = [r for r in wheel.reprofilings if r.location_id == user.location_id]
    at_us = _is_at(wheel, user.location_id)
    current = crud.open_reprofiling(wheel) if at_us else None
    return schemas.RepairWheelOut(
        id=wheel.id,
        serial_number=wheel.serial_number,
        manufacturer=wheel.manufacturer.name,
        diameter=wheel.diameter,
        profile=wheel.profile.code,
        status=wheel.status,
        at_us=at_us,
        state=crud.reprofiling_state(current) if current else ("ceka" if at_us else None),
        sent_at=current.sent_at if current else None,
        total_meters_ground=crud.wheel_total_meters(db, wheel.id),
        reprofilings_here=len([r for r in here if r.result == "hotovo"]),
        # Firma vidí jen svoje reprofilace, ne jiné firmy
        reprofilings=[schemas.WheelReprofilingOut.model_validate(r) for r in here],
    )


def _my_wheels(db: Session, user: models.User) -> list[models.Wheel]:
    """Kotouče, které jsou nebo byly u této firmy."""
    ids_history = select(models.WheelReprofiling.wheel_id).where(models.WheelReprofiling.location_id == user.location_id)
    wheels = db.scalars(
        select(models.Wheel).where(or_(
            models.Wheel.id.in_(ids_history),
            (models.Wheel.status == S.NA_REPROFILACI) & (models.Wheel.location_id == user.location_id),
        ))
    ).all()
    return list(wheels)


@router.get("/summary")
def summary(db: Session = Depends(get_db), user: models.User = Depends(require_repair)):
    """Počty pro dlaždice: čeká na převzetí / v práci / hotovo k odvozu / nelze."""
    active = [w for w in _my_wheels(db, user) if _is_at(w, user.location_id)]
    states = [_to_out(db, w, user).state for w in active]
    return {
        "location": user.location.name if user.location else None,
        "waiting": states.count("ceka"),
        "in_progress": states.count("prevzato"),
        "done": states.count("hotovo"),
        "failed": states.count("nelze"),
    }


@router.get("/wheels", response_model=list[schemas.RepairWheelOut])
def list_wheels(
    scope: str = Query("aktivni", pattern="^(aktivni|historie)$"),
    search: str | None = None,
    db: Session = Depends(get_db),
    user: models.User = Depends(require_repair),
):
    wheels = _my_wheels(db, user)
    if search:
        q = search.strip().lower()
        wheels = [w for w in wheels if q in w.serial_number.lower() or q in w.manufacturer.name.lower()]
    if scope == "aktivni":
        wheels = [w for w in wheels if _is_at(w, user.location_id)]
    else:
        wheels = [w for w in wheels if not _is_at(w, user.location_id)]
    out = [_to_out(db, w, user) for w in wheels]
    order = {"ceka": 0, "prevzato": 1, "nelze": 2, "hotovo": 3}
    return sorted(out, key=lambda o: (order.get(o.state or "", 9), o.sent_at or datetime.min))


@router.get("/wheels/{wheel_id}", response_model=schemas.RepairWheelOut)
def get_wheel(wheel_id: int, db: Session = Depends(get_db), user: models.User = Depends(require_repair)):
    return _to_out(db, _get_wheel(db, wheel_id, user), user)


@router.post("/wheels/{wheel_id}/prevzato", response_model=schemas.RepairWheelOut)
def accept(wheel_id: int, db: Session = Depends(get_db), user: models.User = Depends(require_repair)):
    """Kotouč do firmy dorazil."""
    wheel = _get_wheel(db, wheel_id, user)
    current = _current_here(db, wheel, user)
    if current.result:
        raise HTTPException(400, "Reprofilace už je nahlášená jako dokončená")
    current.accepted_at, current.accepted_by_username = models.utc_now(), user.username
    db.commit()
    db.refresh(wheel)
    return _to_out(db, wheel, user)


@router.post("/wheels/{wheel_id}/hotovo", response_model=schemas.RepairWheelOut)
def done(wheel_id: int, data: schemas.RepairDone, db: Session = Depends(get_db), user: models.User = Depends(require_repair)):
    """Reprofilace hotová — kotouč je připravený k odvozu do AGC. Lze zapsat znovu (např. opravit cenu)."""
    wheel = _get_wheel(db, wheel_id, user)
    current = _current_here(db, wheel, user)
    now = models.utc_now()
    if current.accepted_at is None:  # převzetí se nemuselo zapsat zvlášť
        current.accepted_at, current.accepted_by_username = now, user.username
    current.result, current.finished_at, current.finished_by_username = "hotovo", now, user.username
    current.note = (data.note or "").strip() or None
    current.cost = data.cost
    db.commit()
    db.refresh(wheel)
    return _to_out(db, wheel, user)


@router.post("/wheels/{wheel_id}/nelze", response_model=schemas.RepairWheelOut)
def failed(wheel_id: int, data: schemas.RepairFailed, db: Session = Depends(get_db), user: models.User = Depends(require_repair)):
    """Kotouč nejde reprofilovat (např. málo materiálu) — AGC rozhodne o vyřazení."""
    wheel = _get_wheel(db, wheel_id, user)
    current = _current_here(db, wheel, user)
    now = models.utc_now()
    if current.accepted_at is None:
        current.accepted_at, current.accepted_by_username = now, user.username
    current.result, current.finished_at, current.finished_by_username = "nelze", now, user.username
    current.note = data.reason.strip()
    current.cost = None
    db.commit()
    db.refresh(wheel)
    return _to_out(db, wheel, user)


# ---------- Export do Excelu (jen v části Reprofilace) ----------

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
RESULT = {"hotovo": "Hotovo", "nelze": "Nelze reprofilovat", None: "Rozpracováno"}


def _sheet(wb: Workbook, title: str, columns: list[tuple[str, int, str | None]], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append([c[0] for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    for row in rows:
        ws.append(row)
    for i, (_, width, fmt) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
        if fmt:
            for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell.number_format = fmt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


@router.get("/export.xlsx")
def export(db: Session = Depends(get_db), user: models.User = Depends(require_repair)):
    """Excel: u každého kotouče číslo, kolikrát byl na reprofilaci a u jakých firem (+ cena);
    na druhém listu každá reprofilace této firmy zvlášť."""
    wheels = _my_wheels(db, user)
    rows, summary_rows = [], []
    for w in sorted(wheels, key=lambda x: x.serial_number):
        here = sorted((r for r in w.reprofilings if r.location_id == user.location_id), key=lambda r: r.sent_at)
        for r in here:
            rows.append([
                w.serial_number, w.manufacturer.name, w.diameter, w.profile.code,
                r.sent_at, r.accepted_at, r.finished_at, RESULT.get(r.result, r.result),
                r.cost, r.meters_before, r.note or "", r.returned_at,
            ])
        done_here = [r for r in here if r.result == "hotovo"]
        done_all = sorted((r for r in w.reprofilings if r.result == "hotovo"), key=lambda r: r.sent_at)
        firms: dict[str, int] = {}
        for r in done_all:
            firms[r.location_name or "neuvedeno"] = firms.get(r.location_name or "neuvedeno", 0) + 1
        summary_rows.append([
            w.serial_number, w.manufacturer.name, w.diameter, w.profile.code,
            len(done_all),
            ", ".join(f"{name} ({n}×)" for name, n in firms.items()) or "—",
            len(done_here),
            sum(r.cost or 0 for r in done_here),
            done_all[-1].finished_at if done_all else None,
            "u nás" if _is_at(w, user.location_id) else "vráceno AGC",
        ])

    wb = Workbook()
    wb.remove(wb.worksheets[0])
    dt, czk = "d.m.yyyy h:mm", '# ##0 "Kč"'
    _sheet(wb, "Kotouče", [
        ("Číslo kotouče", 16, None), ("Výrobce", 16, None), ("Průměr (mm)", 11, "0"), ("Profil", 9, None),
        ("Kolikrát na reprofilaci", 20, "0"), ("U jakých firem", 42, None), ("Z toho u nás", 13, "0"),
        ("Cena u nás celkem", 17, czk), ("Poslední reprofilace", 19, dt), ("Kde je teď", 14, None),
    ], summary_rows)
    _sheet(wb, "Reprofilace u nás", [
        ("Číslo kotouče", 16, None), ("Výrobce", 16, None), ("Průměr (mm)", 11, "0"), ("Profil", 9, None),
        ("Odesláno z AGC", 17, dt), ("Převzato", 17, dt), ("Dokončeno", 17, dt), ("Výsledek", 18, None),
        ("Cena reprofilace", 16, czk), ("Nabroušeno před (m)", 18, "0"), ("Poznámka", 40, None), ("Vráceno do AGC", 17, dt),
    ], rows)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"reprofilace-{date.today():%Y-%m-%d}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )
